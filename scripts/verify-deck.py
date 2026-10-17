"""Relecture du deck : structure, textes, images, notes."""

from pathlib import Path

from pptx import Presentation

path = Path(__file__).resolve().parent.parent / "docs" / "Spindown-Codex.pptx"
prs = Presentation(path)

print(f"fichier      : {path.name} ({path.stat().st_size // 1024} Ko)")
print(f"diapositives : {len(prs.slides)}")
print(f"format       : {prs.slide_width.inches:.2f} x {prs.slide_height.inches:.2f} in")

problems = []
for index, slide in enumerate(prs.slides, start=1):
    texts = []
    pictures = 0
    tables = 0
    for shape in slide.shapes:
        if shape.shape_type == 13:
            pictures += 1
        if shape.has_table:
            tables += 1
        if shape.has_text_frame and shape.text_frame.text.strip():
            texts.append(shape.text_frame.text.strip())

    notes = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
    title = texts[1] if len(texts) > 1 else (texts[0] if texts else "")
    print(f"  {index:02d}  images={pictures} tables={tables} notes={len(notes):>3}  {title[:60]}")

    if not texts:
        problems.append(f"diapo {index} sans texte")
    for shape in slide.shapes:
        if shape.has_text_frame:
            box = shape.left + shape.width
            if box > prs.slide_width:
                problems.append(f"diapo {index} : debordement horizontal")

print("\nproblemes :", problems if problems else "aucun")
